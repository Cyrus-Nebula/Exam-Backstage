from fastapi import FastAPI, Response
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from typing import List, Optional
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from io import BytesIO
import httpx
import re
import html as html_lib
import base64

app = FastAPI(title="Personal Exam Export")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


class Question(BaseModel):
    id: Optional[str] = ""
    html: str = ""
    text: str = ""
    answer: str = ""
    analysis: str = ""
    images: List[str] = []


class Section(BaseModel):
    title: str = "题目"
    questions: List[Question] = []


class Paper(BaseModel):
    title: str = "未命名试卷"
    site: str = ""
    url: str = ""
    time: str = ""
    sections: List[Section] = []


class ExportOptions(BaseModel):
    includeAnswer: bool = True
    answerAtEnd: bool = False


class ExportRequest(BaseModel):
    paper: Paper
    options: ExportOptions = ExportOptions()


def strip_html(s: str) -> str:
    if not s:
        return ""
    s = re.sub(r"<br\s*/?>", "\n", s, flags=re.I)
    s = re.sub(r"</p>", "\n", s, flags=re.I)
    s = re.sub(r"<[^>]+>", "", s)
    s = html_lib.unescape(s)
    return re.sub(r"\n{3,}", "\n\n", s).strip()


def set_run_font(run, size=12, bold=False, east_asia="宋体"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    try:
        run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    except Exception:
        pass


async def download_image(url: str, client: httpx.AsyncClient):
    if not url:
        return None
    try:
        if url.startswith("data:image"):
            _header, b64 = url.split(",", 1)
            data = base64.b64decode(b64)
            return BytesIO(data)
        r = await client.get(url, timeout=20.0, follow_redirects=True)
        if r.status_code == 200 and r.content:
            return BytesIO(r.content)
    except Exception as e:
        print("img fail", url, e)
    return None


async def build_docx(paper: Paper, opts: ExportOptions) -> bytes:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run(paper.title or "试卷")
    set_run_font(run, size=18, bold=True)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mrun = meta.add_run(f"{paper.site}  ·  {paper.time}")
    set_run_font(mrun, size=10)
    mrun.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    doc.add_paragraph()
    answer_bucket = []

    async with httpx.AsyncClient(
        headers={"User-Agent": "Mozilla/5.0 (compatible; PersonalExport/1.0)"}
    ) as client:
        for sec in paper.sections:
            h = doc.add_paragraph()
            hr = h.add_run(sec.title)
            set_run_font(hr, size=14, bold=True)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(8)

            for i, q in enumerate(sec.questions, 1):
                body_text = strip_html(q.html) or q.text or ""
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(4)
                num_run = p.add_run(f"{i}. ")
                set_run_font(num_run, size=12, bold=True)
                if body_text:
                    tr = p.add_run(body_text)
                    set_run_font(tr, size=12)

                for img_url in (q.images or [])[:8]:
                    bio = await download_image(img_url, client)
                    if bio:
                        try:
                            ip = doc.add_paragraph()
                            ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            pr = ip.add_run()
                            pr.add_picture(bio, width=Inches(3.8))
                        except Exception as e:
                            print("pic insert fail", e)

                if opts.includeAnswer and not opts.answerAtEnd:
                    if q.answer:
                        ap = doc.add_paragraph()
                        ar = ap.add_run(f"答案：{q.answer}")
                        set_run_font(ar, size=11)
                        ar.font.color.rgb = RGBColor(0x33, 0x55, 0x99)
                    if q.analysis:
                        ap2 = doc.add_paragraph()
                        ar2 = ap2.add_run(f"解析：{q.analysis}")
                        set_run_font(ar2, size=11)
                        ar2.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

                if opts.includeAnswer and opts.answerAtEnd:
                    answer_bucket.append((sec.title, i, q.answer, q.analysis))

    if opts.includeAnswer and opts.answerAtEnd and answer_bucket:
        doc.add_page_break()
        ah = doc.add_paragraph()
        ahr = ah.add_run("参考答案与解析")
        set_run_font(ahr, size=16, bold=True)
        for sec_title, num, ans, ana in answer_bucket:
            p = doc.add_paragraph()
            pr = p.add_run(f"{sec_title} 第{num}题")
            set_run_font(pr, size=12, bold=True)
            if ans:
                p2 = doc.add_paragraph()
                r2 = p2.add_run(f"答案：{ans}")
                set_run_font(r2, size=11)
            if ana:
                p3 = doc.add_paragraph()
                r3 = p3.add_run(f"解析：{ana}")
                set_run_font(r3, size=11)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


@app.get("/")
def root():
    return {"ok": True, "msg": "Personal Exam Export Server running"}


@app.post("/export")
async def export(req: ExportRequest):
    data = await build_docx(req.paper, req.options)
    filename = (req.paper.title or "试卷").replace("/", "_")[:40] + ".docx"
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Access-Control-Allow-Origin": "*",
        },
    )
